"""Offline workbook inspection/scoring. Never expose reference files to agent tools.

Value normalization follows SpreadsheetBench revision 49b73a94775fb489063f60ca1865e3a650079a79.
This is a local comparison adapter, not the official three-fixture evaluation CLI.
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.cell import range_boundaries

MAX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_CELLS = 100_000


def inspect_workbook(path: Path) -> dict:
    if path.is_symlink():
        raise ValueError("symlink workbook rejected")
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
            raise ValueError("expanded workbook too large")
        names = archive.namelist()
        if any("vbaproject" in n.lower() or "externallinks/" in n.lower() for n in names):
            raise ValueError("macro or external workbook link rejected")
    book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = []
        for sheet in book:
            if sheet.max_row * sheet.max_column > MAX_CELLS:
                raise ValueError("worksheet dimensions exceed inspection limit")
            formulas = sum(c.data_type == "f" for row in sheet for c in row)
            sheets.append({"name": sheet.title, "rows": sheet.max_row,
                           "columns": sheet.max_column, "formulas": formulas})
        return {"sheets": sheets, "requires_recalculation": any(s["formulas"] for s in sheets)}
    finally:
        book.close()


def normalized(value):
    if isinstance(value, dt.datetime):
        return round((value - dt.datetime(1899, 12, 30)).total_seconds() / 86400, 0)
    if isinstance(value, dt.time):
        return str(value)[:-3]
    if isinstance(value, (str, int, float)):
        try:
            return round(float(value), 2)
        except ValueError:
            return value
    return value


def same_value(left, right):
    left, right = normalized(left), normalized(right)
    if left in (None, "") and right in (None, ""):
        return True
    return type(left) is type(right) and left == right


def answer_ranges(position: str, default_sheet: str):
    for entry in position.split(","):
        sheet, marker, cells = entry.partition("!")
        if not marker:
            sheet, cells = default_sheet, sheet
        sheet, cells = sheet.strip("'"), cells.strip("'")
        bounds = range_boundaries(cells)
        if None in bounds or (bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1) > MAX_CELLS:
            raise ValueError("invalid or oversized answer range")
        yield sheet, bounds


def compare_workbooks(reference: Path, output: Path, position: str) -> dict:
    # A broken reference/scoring setup is infrastructure, not an agent failure.
    try:
        reference_info = inspect_workbook(reference)
        expected = openpyxl.load_workbook(reference, data_only=True, keep_links=False)
        ranges = list(answer_ranges(position, expected.sheetnames[0]))
        if not ranges or any(name not in expected for name, _ in ranges):
            raise ValueError("unknown reference worksheet")
    except Exception as exc:
        return {"status": "INFRA_ERROR", "code": "ORACLE_EXECUTION_ERROR", "detail": str(exc)}
    try:
        try:
            output_info = inspect_workbook(output)
            # Formula results must not silently compare empty caches as correct.
            if reference_info["requires_recalculation"] or output_info["requires_recalculation"]:
                return {"status": "INFRA_ERROR", "code": "ORACLE_EXECUTION_ERROR",
                        "detail": "formula recalculation evidence required; not implemented yet"}
            actual = openpyxl.load_workbook(output, data_only=True, keep_links=False)
        except Exception as exc:
            return {"status": "TASK_FAIL", "code": "ORACLE_ASSERTION_FAILED", "detail": str(exc)}
        try:
            differences = []
            checked = 0
            for name, (c1, r1, c2, r2) in ranges:
                if name not in actual:
                    differences.append({"sheet": name, "reason": "missing worksheet"})
                    continue
                for row in range(r1, r2 + 1):
                    for col in range(c1, c2 + 1):
                        checked += 1
                        cell = expected[name].cell(row, col)
                        if not same_value(cell.value, actual[name].cell(row, col).value):
                            differences.append({"sheet": name, "cell": cell.coordinate})
            integrity = expected.sheetnames == actual.sheetnames
            for name in expected.sheetnames:
                if name not in actual:
                    integrity = False
                    continue
                left, right = expected[name], actual[name]
                for row in range(1, max(left.max_row, right.max_row) + 1):
                    for col in range(1, max(left.max_column, right.max_column) + 1):
                        if not same_value(left.cell(row, col).value, right.cell(row, col).value):
                            integrity = False
            return {"status": "TASK_FAIL" if differences else "TASK_PASS",
                    "comparison": "spreadsheetbench-value-normalization-local-adapter-v1",
                    "checked_cells": checked, "difference_count": len(differences),
                    "differences": differences[:100],
                    "separate_audit": {"all_workbook_values_match_reference": integrity,
                                       "affects_upstream_score": False},
                    "scope": "specified cells only; formatting and extra rows not scored"}
        finally:
            actual.close()
    finally:
        expected.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("reference", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("answer_position")
    args = parser.parse_args()
    print(json.dumps(compare_workbooks(args.reference, args.output, args.answer_position), indent=2))
