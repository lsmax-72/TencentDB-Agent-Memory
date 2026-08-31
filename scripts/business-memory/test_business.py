import asyncio
import datetime as dt
import io
import json
import os
import tarfile
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import openpyxl

import preflight
from sandbox import IMAGE, PythonSandbox
from workbook_oracle import answer_ranges, compare_workbooks, inspect_workbook, same_value


def book(path, value="INV-010", second_sheet=True):
    workbook = openpyxl.Workbook()
    workbook.active.title = "Ledger"
    workbook.active.append(["Invoice", "Amount"])
    workbook.active.append([value, 25.5])
    if second_sheet:
        workbook.create_sheet("Notes").append(["preserve me"])
    workbook.save(path)


class OracleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.reference, self.output = self.root / "golden.xlsx", self.root / "result.xlsx"
        book(self.reference)
        book(self.output)

    def compare(self):
        return compare_workbooks(self.reference, self.output, "'Ledger!'A1:B2")

    def test_correct_and_wrong_cell(self):
        self.assertEqual(self.compare()["status"], "TASK_PASS")
        book(self.output, "WRONG")
        result = self.compare()
        self.assertEqual(result["status"], "TASK_FAIL")
        self.assertEqual(result["differences"], [{"sheet": "Ledger", "cell": "A2"}])

    def test_missing_and_corrupted_output_are_task_failures(self):
        self.output.unlink()
        self.assertEqual(self.compare()["status"], "TASK_FAIL")
        self.output.write_bytes(b"not an xlsx")
        self.assertEqual(self.compare()["status"], "TASK_FAIL")

    def test_missing_reference_and_invalid_range_are_infra(self):
        self.assertEqual(compare_workbooks(self.reference, self.output, "Missing!A1")["status"], "INFRA_ERROR")
        self.reference.unlink()
        self.assertEqual(self.compare()["status"], "INFRA_ERROR")

    def test_symlink_output_is_rejected(self):
        self.output.unlink()
        self.output.symlink_to(self.reference)
        self.assertEqual(self.compare()["status"], "TASK_FAIL")

    def test_extra_rows_and_other_sheet_are_separate_audit(self):
        workbook = openpyxl.load_workbook(self.output)
        workbook["Ledger"].append(["unrequested", 99])
        workbook["Notes"]["A1"] = "changed"
        workbook.save(self.output)
        result = self.compare()
        self.assertEqual(result["status"], "TASK_PASS")
        self.assertFalse(result["separate_audit"]["all_workbook_values_match_reference"])

    def test_formulas_require_real_recalculation(self):
        book(self.output, "=1+1")
        self.assertTrue(inspect_workbook(self.output)["requires_recalculation"])
        self.assertEqual(self.compare()["status"], "INFRA_ERROR")

    def test_normalization_matches_upstream_semantics(self):
        self.assertTrue(same_value("25.500", 25.5))
        self.assertTrue(same_value(None, ""))
        self.assertTrue(same_value(dt.datetime(1900, 1, 1), 2))
        self.assertTrue(same_value(dt.time(12, 30), "12:30"))
        self.assertFalse(same_value("INV-010", "INV-10"))

    def test_quoted_ranges(self):
        self.assertEqual(list(answer_ranges("'A!'A1:B2,'B'!C1", "X")),
                         [("A", (1, 1, 2, 2)), ("B", (3, 1, 3, 1))])


class FixtureTests(unittest.TestCase):
    def test_archive_checksum_rejection(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "archive"
            path.write_bytes(b"untrusted")
            with self.assertRaisesRegex(ValueError, "unrecognized archive"):
                preflight.prepare(path, Path(tmp) / "new")
            self.assertFalse((Path(tmp) / "new").exists())

    def test_verified_layout_and_never_overwrite(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            book(root / "sheet.xlsx")
            archive = root / "fixture.tar.gz"
            with tarfile.open(archive, "w:gz") as tar:
                for name, payload in [
                    ("data/dataset.json", json.dumps([{"id": "141-20"}]).encode()),
                    ("data/spreadsheet/141-20/1_141-20_init.xlsx", (root / "sheet.xlsx").read_bytes()),
                    ("data/spreadsheet/141-20/1_141-20_golden.xlsx", (root / "sheet.xlsx").read_bytes()),
                ]:
                    item = tarfile.TarInfo(name)
                    item.size = len(payload)
                    tar.addfile(item, io.BytesIO(payload))
            with patch.object(preflight, "ARCHIVE_SHA256", preflight.digest(archive)):
                result = preflight.prepare(archive, root / "prepared")
                self.assertEqual([f["role"] for f in result["files"]], ["input", "reference"])
                self.assertEqual(result['data_role'], 'development_not_heldout')
                confirmation = preflight.prepare(archive, root / 'confirmation', data_role='reserved_policy_confirmation')
                self.assertEqual(confirmation['data_role'], 'reserved_policy_confirmation')
                with self.assertRaises(FileExistsError):
                    preflight.prepare(archive, root / "prepared")


class ToolMappingTests(unittest.TestCase):
    def test_sdk_registration_and_error_mapping_without_model(self):
        from nanobot.agent.tools.registry import ToolRegistry
        from nanobot_tool import register_spreadsheet_only
        result = {"ok": False, "output": "Permission denied"}
        class FakeSandbox:
            def run(self, code):
                self.code = code
                return result
        sandbox = FakeSandbox()
        registry = ToolRegistry()
        bot = SimpleNamespace(_loop=SimpleNamespace(tools=registry))
        register_spreadsheet_only(bot, sandbox)
        self.assertEqual(registry.tool_names, ["spreadsheet_python"])
        tool = registry.get("spreadsheet_python")
        output = asyncio.run(tool.execute("print(1)"))
        self.assertTrue(output.is_error)
        self.assertEqual(sandbox.code, "print(1)")
        self.assertEqual(json.loads(output), result)

    def test_protocol_requires_explicit_decision_and_preserves_model(self):
        from runner_contract import build_config, validate_protocol
        protocol = json.loads(Path(__file__).with_name("protocol-smoke-v1.json").read_text())
        protocol.update(status="REVIEW_REQUIRED", approval_ref=None)
        with self.assertRaisesRegex(ValueError, "REVIEW_REQUIRED"):
            validate_protocol(protocol)
        protocol.update(status="APPROVED_BY_USER", approval_ref="test-only-decision")
        config = build_config(protocol, proxy_url="http://127.0.0.1:19296/proxy/business-test/v1",
                              user_key="test-key", identity={"team_id": "t", "agent_id": "a",
                              "task_id": "job", "session_id": "session-unique"})
        self.assertEqual(config["agents"]["defaults"]["provider"], "vllm")
        self.assertEqual(config["agents"]["defaults"]["fallbackModels"], [])
        self.assertEqual(config["providers"]["vllm"]["extraHeaders"]["x-session-id"], "session-unique")
        protocol["temperature"] = 0.1
        with self.assertRaisesRegex(ValueError, "RUNSPEC_MISMATCH"):
            validate_protocol(protocol)

    def test_old_proxy_and_incomplete_usage_are_rejected(self):
        from runner_contract import build_config, usage_summary
        protocol = json.loads(Path(__file__).with_name("protocol-smoke-v1.json").read_text())
        protocol.update(status="APPROVED_BY_USER", approval_ref="test-only-decision")
        with self.assertRaisesRegex(ValueError, "historical proxy"):
            build_config(protocol, proxy_url="http://127.0.0.1:19096/proxy/business-test/v1",
                         user_key="test-key", identity={})
        with self.assertRaisesRegex(ValueError, "TELEMETRY_INCOMPLETE"):
            usage_summary({"actual_model": "qwen3.8-27b", "usage": {}}, [], "qwen3.8-27b")
        raw = {"actual_model": "qwen3.8-27b", "model_calls": 1,
               "usage": {"prompt_tokens": 10, "completion_tokens": 5, "total_tokens": 15}}
        result = usage_summary(raw, [{"name": "spreadsheet_python", "outcome": "SUCCEEDED", "result": "ok"}],
                               "qwen3.8-27b")
        self.assertEqual(result["tool_calls"], 1)


@unittest.skipUnless(os.environ.get("BUSINESS_DOCKER_TESTS") == "1", "explicit local Docker test opt-in")
class RealSandboxTests(unittest.TestCase):
    def setUp(self):
        import et_xmlfile
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.inputs, self.outputs = self.root / "inputs", self.root / "outputs"
        self.inputs.mkdir(mode=0o755)
        self.outputs.mkdir(mode=0o777)
        self.outputs.chmod(0o777)
        book(self.inputs / "input.xlsx")
        (self.inputs / "input.xlsx").chmod(0o644)
        self.sandbox = PythonSandbox(self.inputs, self.outputs, Path(et_xmlfile.__file__).parent.parent)

    def test_real_xlsx_read_write(self):
        result = self.sandbox.run("import openpyxl; w=openpyxl.load_workbook('/inputs/input.xlsx'); "
                                  "w['Ledger'].delete_rows(2); w.save('/outputs/output.xlsx'); print('done')")
        self.assertTrue(result["ok"], result)
        self.assertEqual(inspect_workbook(self.outputs / "output.xlsx")["sheets"][0]["rows"], 1)

    def test_network_host_credentials_and_input_write_denied(self):
        code = """
import os, socket
from pathlib import Path
assert os.getuid() != 0
assert not Path('/var/run/docker.sock').exists()
assert not Path('/Users/lsmax/.nanobot/config.json').exists()
assert not Path('/reference').exists()
for dest in ['/inputs/input.xlsx', '/etc/business-test']:
    try:
        open(dest, 'wb').write(b'bad')
    except OSError:
        pass
    else:
        raise AssertionError('unexpected write')
try:
    socket.create_connection(('1.1.1.1', 443), timeout=1)
except OSError:
    pass
else:
    raise AssertionError('unexpected network')
print('denials verified')
"""
        result = self.sandbox.run(code)
        self.assertTrue(result["ok"], result)

    def test_timeout_and_large_stdout(self):
        result = self.sandbox.run("import time; time.sleep(30)", timeout_seconds=2)
        self.assertEqual(result["stop_reason"], "TOOL_TIMEOUT")
        result = self.sandbox.run("print('x'*200000)")
        self.assertEqual(result["stop_reason"], "TOOL_OUTPUT_LIMIT")

    def test_symlink_output_rejected(self):
        result = self.sandbox.run("import os; os.symlink('/etc/passwd', '/outputs/link')")
        self.assertEqual(result["stop_reason"], "TOOL_POLICY_VIOLATION")

    def test_resource_limits_and_mutable_image_rejection(self):
        result = self.sandbox.run("import resource; assert resource.getrlimit(resource.RLIMIT_FSIZE)[0] == 16777216; "
                                  "b=bytearray(900*1024*1024)")
        self.assertFalse(result["ok"])
        with self.assertRaisesRegex(ValueError, "immutable"):
            PythonSandbox(self.inputs, self.outputs, self.sandbox.packages, "python:latest")


if __name__ == "__main__":
    unittest.main()
